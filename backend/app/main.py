import os
from pathlib import Path

from dotenv import load_dotenv
from fastapi import FastAPI, Depends
from fastapi.middleware.cors import CORSMiddleware
from app.customers.routes import router as customers_router
from app.services.routes import router as services_router
from app.inventory.routes import materials_router, inventory_router
from app.reminders.routes import router as reminders_router
from app.reports.routes import router as reports_router
from app.auth import auth_router, require_auth
from app.excel_store import DATA_DIR, CUSTOMERS_FILE, MATERIALS_FILE
from openpyxl import Workbook, load_workbook

load_dotenv()

app = FastAPI(title="Maruti Enterprises API")

# ---------------------------------------------------------------------------
# Security Headers Middleware (Pure ASGI to prevent request streaming deadlocks)
# ---------------------------------------------------------------------------

class SecurityHeadersMiddleware:
    def __init__(self, app):
        self.app = app

    async def __call__(self, scope, receive, send):
        if scope["type"] != "http":
            await self.app(scope, receive, send)
            return

        async def send_wrapper(message):
            if message["type"] == "http.response.start":
                headers = list(message.get("headers", []))
                headers.append((b"x-content-type-options", b"nosniff"))
                headers.append((b"x-frame-options", b"DENY"))
                headers.append((b"referrer-policy", b"strict-origin-when-cross-origin"))
                message["headers"] = headers
            await send(message)

        await self.app(scope, receive, send_wrapper)


app.add_middleware(SecurityHeadersMiddleware)



# ---------------------------------------------------------------------------
# CORS Configuration
# ---------------------------------------------------------------------------

allowed_origins_env = os.getenv("ALLOWED_ORIGINS", "http://localhost:3000,http://127.0.0.1:3000")
allowed_origins = [origin.strip() for origin in allowed_origins_env.split(",") if origin.strip()]

# Support all local loopback addresses and LAN private subnets (for mobile & desktop dev)
local_network_origin_regex = (
    r"^https?://(localhost|127\.0\.0\.1|0\.0\.0\.0|"
    r"192\.168\.\d+\.\d+|10\.\d+\.\d+\.\d+|172\.(1[6-9]|2\d|3[0-1])\.\d+\.\d+)(:\d+)?$"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=allowed_origins,
    allow_origin_regex=local_network_origin_regex,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)



# ---------------------------------------------------------------------------
# Excel column definitions
# ---------------------------------------------------------------------------

CUSTOMERS_SHEETS = {
    "Customers": [
        "customer_id",
        "name",
        "mobile",
        "alternate_mobile",
        "address",
        "area",
        "ro_brand",
        "ro_model",
        "installation_date",
        "technician",
        "notes",
        "created_at",
    ],
    "Services": [
        "service_id",
        "customer_id",
        "service_date",
        "service_type",
        "technician",
        "complaint",
        "work_done",
        "next_service_date",
        "labour_charge",
        "material_charge",
        "total_amount",
        "payment_status",
        "notes",
    ],
    "Service_Materials": [
        "service_id",
        "material_id",
        "quantity",
    ],
}

MATERIALS_SHEETS = {
    "Materials": [
        "material_id",
        "material_name",
        "category",
        "unit",
        "current_stock",
        "minimum_stock",
        "purchase_price",
        "selling_price",
        "supplier",
    ],
    "Stock_Transactions": [
        "transaction_id",
        "material_id",
        "transaction_type",
        "quantity",
        "date",
        "reference_id",
        "notes",
    ],
}


# ---------------------------------------------------------------------------
# Excel initialization helpers
# ---------------------------------------------------------------------------

def _create_excel_file(file_path: Path, sheets: dict[str, list[str]]) -> None:
    """Create an Excel workbook with the given sheets and column headers."""
    wb = Workbook()
    # Remove the default sheet created by openpyxl
    default_sheet = wb.active
    wb.remove(default_sheet)

    for sheet_name, columns in sheets.items():
        ws = wb.create_sheet(title=sheet_name)
        ws.append(columns)

    wb.save(file_path)


def initialize_excel_files() -> None:
    """Ensure both Excel data files exist with the correct sheets and headers."""
    DATA_DIR.mkdir(parents=True, exist_ok=True)

    default_data_dir = Path(__file__).resolve().parent.parent / "data"
    default_customers = default_data_dir / "customers.xlsx"
    default_materials = default_data_dir / "materials.xlsx"

    if not CUSTOMERS_FILE.exists():
        if default_customers.exists() and default_customers.resolve() != CUSTOMERS_FILE.resolve():
            import shutil
            shutil.copy2(default_customers, CUSTOMERS_FILE)
        else:
            _create_excel_file(CUSTOMERS_FILE, CUSTOMERS_SHEETS)

    if not MATERIALS_FILE.exists():
        if default_materials.exists() and default_materials.resolve() != MATERIALS_FILE.resolve():
            import shutil
            shutil.copy2(default_materials, MATERIALS_FILE)
        else:
            _create_excel_file(MATERIALS_FILE, MATERIALS_SHEETS)


# Run initialization at startup
initialize_excel_files()


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------

@app.get("/health")
def health_check():
    """Public health-check endpoint."""
    return {"status": "ok"}


# Authentication routes (login, logout, me)
app.include_router(auth_router)

# Protected Business Management routes
app.include_router(customers_router, dependencies=[Depends(require_auth)])
app.include_router(services_router, dependencies=[Depends(require_auth)])
app.include_router(materials_router, dependencies=[Depends(require_auth)])
app.include_router(inventory_router, dependencies=[Depends(require_auth)])
app.include_router(reminders_router, dependencies=[Depends(require_auth)])
app.include_router(reports_router, dependencies=[Depends(require_auth)])

